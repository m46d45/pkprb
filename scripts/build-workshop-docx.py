#!/usr/bin/env python3
"""Generate the KoNTekS 20 validation workshop runbook (DOCX)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INK = RGBColor(0x1F, 0x1A, 0x16)
MUTED = RGBColor(0x6B, 0x63, 0x58)
TEAL = RGBColor(0x3D, 0x7F, 0x92)
HEADER_BG = "1F1A16"
ROW_ALT = "F7F3EC"
FORMULA_BG = "EDE7DC"

ROOT = Path(__file__).resolve().parents[1]
OUT_PUBLIC = ROOT / "public" / "Panduan-Acara-Validasi-PKPRB.docx"
OUT_DOCS = ROOT / "docs" / "Panduan-Acara-Validasi-PKPRB.docx"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=INK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, color="D9D1C5"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(
        "PKPRB  ·  Panduan acara validasi  ·  KoNTekS 20  ·  pkprb.vercel.app"
    )
    set_run_font(run, size=8, color=MUTED)
    p.add_run("    ")
    run2 = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run2._r.append(fld)
    run3 = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run3._r.append(instr)
    run4 = p.add_run()
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run4._r.append(fld2)
    for r in (run2, run3, run4):
        set_run_font(r, size=8, color=MUTED)


def p_style(paragraph, space_after=8, space_before=0):
    pf = paragraph.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.15
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = INK
        run.font.name = "Cambria"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria")
    p_style(h, space_after=8, space_before=16 if level == 1 else 12)
    return h


def body(doc, text, *, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    p_style(p)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        p_style(p, space_after=4)
        if isinstance(item, str):
            run = p.add_run(item)
            set_run_font(run)
        else:
            for text, bold, italic in item:
                run = p.add_run(text)
                set_run_font(run, bold=bold, italic=italic)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.autofit = True
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        shade(cell, HEADER_BG)
        set_cell_border(cell, HEADER_BG)
        p = cell.paragraphs[0]
        p_style(p, space_after=0)
        run = p.add_run(h)
        set_run_font(run, size=9, bold=True, color=RGBColor(0xF3, 0xEE, 0xE4))
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = t.rows[r_i + 1].cells[c_i]
            if r_i % 2 == 1:
                shade(cell, ROW_ALT)
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p_style(p, space_after=0)
            run = p.add_run(str(val))
            set_run_font(run, size=9)
    doc.add_paragraph()
    return t


def callout(doc, title, body_text, fill="E8D9D0", border="C45C48"):
    t = doc.add_table(rows=1, cols=1)
    cell = t.cell(0, 0)
    shade(cell, fill)
    set_cell_border(cell, border)
    p = cell.paragraphs[0]
    p_style(p, space_after=4)
    run = p.add_run(title)
    set_run_font(run, size=11, bold=True)
    p2 = cell.add_paragraph()
    p_style(p2, space_after=0)
    run2 = p2.add_run(body_text)
    set_run_font(run2, size=10)
    doc.add_paragraph()


def note(doc, title, body_text):
    callout(doc, title, body_text, fill=FORMULA_BG, border="C9C0B3")


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    add_footer(section)

    kicker = doc.add_paragraph()
    p_style(kicker, space_after=2)
    run = kicker.add_run(
        "KONTEKS 20  ·  23 OKTOBER 2026  ·  P2MI MULTIDISIPLIN FTSL ITB"
    )
    set_run_font(run, size=9, bold=True, color=TEAL)

    title = doc.add_paragraph()
    p_style(title, space_after=4, space_before=0)
    run = title.add_run("Panduan acara validasi PKPRB")
    set_run_font(run, name="Cambria", size=26, bold=True)

    sub = doc.add_paragraph()
    p_style(sub, space_after=12)
    run = sub.add_run(
        "Skenario fasilitator untuk workshop hibrid — termasuk peserta dadakan"
    )
    set_run_font(run, name="Cambria", size=14, italic=True, color=MUTED)

    body(
        doc,
        "Dokumen ini untuk fasilitator (Anda), bukan untuk dibagikan ke peserta. "
        "Versi peta yang divalidasi: https://pkprb.vercel.app — IRBI 2025, "
        "form Kirim isian, hasil tersimpan di Neon.",
    )

    callout(
        doc,
        "Satu kalimat pembuka.",
        "“Isi form di HP atau laptop, tekan Kirim isian. Saya yang unduh hasilnya nanti. "
        "Tidak perlu unduh file. Sel yang dikosongkan artinya Bapak/Ibu menyetujui nilai bawaan.”",
    )

    heading(doc, "1. Apa yang divalidasi")
    body(
        doc,
        "Peserta menilai bobot dan asumsi peta — bukan daftar perguruan tinggi, "
        "bukan kesiapan BPBD, bukan mutu lulusan di lapangan.",
    )
    bullets(
        doc,
        [
            "Cara merangkum skor kabupaten IRBI menjadi angka provinsi (A1) dan nilai 0 (A2).",
            "Bobot jenjang, akreditasi, IABEE, pusat studi, kepakaran (PkM), spillover.",
            "Matriks disiplin × jenis bahaya: boleh isi hanya baris sesuai keahlian.",
            "Catatan terbuka: satu hal yang harus diubah / tidak diubah.",
        ],
    )
    note(
        doc,
        "Bukan peta kesiapan daerah.",
        "Filter Komposit = indeks resmi IRBI 2025 tingkat provinsi. Filter gempa, tsunami, "
        "banjir, dan seterusnya = rata-rata skor kabupaten. Angka 0 = tidak ada baris di "
        "tabel ancaman, bukan sertifikat aman.",
    )

    heading(doc, "2. Peran")
    table(
        doc,
        ["Siapa", "Yang dilakukan", "Yang tidak dilakukan"],
        [
            [
                "Peserta / validator",
                "Buka peta, isi form, tekan Kirim isian",
                "Tidak mengunduh JSON, tidak geser peta bersama",
            ],
            [
                "Fasilitator (Anda)",
                "Screen share peta, pimpin diskusi, unduh JSON kapan saja",
                "Jangan share screen tab hasil / token",
            ],
            [
                "Peserta dadakan",
                "Ikut dari menit berjalan; QR form tetap sama",
                "Tidak perlu undangan pra-kerja",
            ],
        ],
    )

    heading(doc, "3. Tautan")
    table(
        doc,
        ["Keperluan", "URL"],
        [
            ["Peta (layar utama)", "https://pkprb.vercel.app"],
            ["Metodologi", "https://pkprb.vercel.app/metodologi"],
            ["Form validator", "https://pkprb.vercel.app/kuesioner.html"],
            [
                "Hasil (hanya fasilitator)",
                "https://pkprb.vercel.app/validasi/hasil",
            ],
        ],
    )
    callout(
        doc,
        "Token fasilitator — jangan diproyeksikan.",
        "Buka /validasi/hasil di tab terpisah. Token: 12345678. "
        "Tampilkan daftar, lalu Unduh semua JSON setelah acara atau di jeda. "
        "Jangan tampilkan token di slide atau screen share.",
    )

    heading(doc, "4. Checklist")
    heading(doc, "H−1 atau pagi hari", level=2)
    bullets(
        doc,
        [
            "Buka peta, form, dan /validasi/hasil dari HP dan laptop; kirim satu isian uji, pastikan muncul di daftar.",
            "Siapkan QR ke https://pkprb.vercel.app/kuesioner.html (cetak atau tampilkan di slide cadangan).",
            "Tulis URL pendek di papan: pkprb.vercel.app",
            "Charger, pointer, dan cadangan hotspot.",
            "Notes atau Sheets kosong untuk keputusan A1–E2 (bukan di layar peserta).",
        ],
    )
    heading(doc, "H−15 menit", level=2)
    bullets(
        doc,
        [
            "Laptop: tab 1 peta, tab 2 metodologi, tab 3 form, tab 4 hasil (jangan dishare).",
            "Uji screen share hanya tab peta.",
            "Cek jumlah jawaban yang sudah ada di /validasi/hasil (catat angka awal).",
            "Mode peta awal: Keselarasan + filter Gempabumi.",
        ],
    )

    heading(doc, "5. Agenda (90–100 menit)")
    table(
        doc,
        ["Waktu", "Blok", "Yang terjadi"],
        [
            [
                "0–10",
                "Orientasi",
                "PKPRB bukan peta kesiapan. Tunjuk Risiko → Pendidikan → Keselarasan. Satu contoh Senjang.",
            ],
            [
                "10–20",
                "Cara mengisi",
                "QR form. Kosong = setuju default. Matriks: 1–2 baris keahlian. Tombol: Kirim isian.",
            ],
            [
                "20–45",
                "Keputusan struktural",
                "Vote / angkat tangan: A1, A2, D1, E1–E2, C1. Geser 1–2 slider di peta.",
            ],
            [
                "45–75",
                "Isi form + peta hidup",
                "Peserta mengisi. Anda keliling / pantau remote. Geser 3–5 bobot kritis. Cek daftar hasil diam-diam.",
            ],
            [
                "75–90",
                "Penutup",
                "Yang belum: Kirim isian. Ulangi izin ringkas vs nama. Satu hal diubah / tidak diubah.",
            ],
            [
                "Pasca",
                "Fasilitator",
                "Unduh semua JSON. Rekap. Usulan ubah src/lib/weights.ts jika ada konsensus.",
            ],
        ],
    )

    heading(doc, "6. Naskah singkat per blok")
    heading(doc, "Orientasi", level=2)
    body(
        doc,
        "“Peta ini menampilkan keselarasan antara risiko IRBI dan dukungan pendidikan tinggi. "
        "Bukan peta kesiapan BPBD. Komposit memakai angka resmi provinsi. Filter gempa dan "
        "sejenisnya memakai rata-rata kabupaten — nol artinya tidak ada baris, bukan daerah aman.”",
    )
    heading(doc, "Cara mengisi", level=2)
    body(
        doc,
        "“Scan QR. Identitas boleh inisial. Bagian yang di luar keahlian boleh dikosongkan. "
        "Matriks: isi hanya sel yang menurut Bapak/Ibu salah. Setelah centang persetujuan, "
        "tekan Kirim isian. Jangan cari tombol unduh.”",
    )
    heading(doc, "Keputusan struktural", level=2)
    body(
        doc,
        "Bacakan opsi, minta angkat tangan, catat mayoritas. Jika terpecah, tulis ‘belum sepakat’ "
        "dan lanjut. Jangan memaksa angka di form semua orang.",
    )
    heading(doc, "Penutup", level=2)
    body(
        doc,
        "“Terima kasih. Jawaban akan diringkas untuk menyetel bobot bawaan. Nama per orang "
        "tidak dikutip kecuali izin institusi dicentang.”",
    )

    heading(doc, "7. Butir diskusi (urut prioritas)")
    table(
        doc,
        ["Kode", "Pertanyaan", "Opsi di form"],
        [
            ["A1", "Agregasi kabupaten → provinsi", "Rata-rata / maksimum / tertimbang penduduk"],
            ["A2", "Provinsi tanpa baris IRBI", "Tetap 0 / isi ahli / sembunyikan filter"],
            ["D1", "Kepakaran = bendera PkM pusat", "Setuju / harus ada daftar kegiatan / hapus"],
            ["E1", "Spillover antarprovinsi", "Tahan / kecilkan / besarkan / hapus"],
            ["E2", "Cara membagi spillover", "Rata sepulau–nasional / tetangga / penerima"],
            ["C1", "IABEE = Internasional, tidak ditumpuk", "Setuju / terlalu tinggi / boleh ditumpuk"],
            ["B1", "S3 vs S1", "Ya lebih tinggi / tidak / bergantung disiplin"],
            ["F1", "Kesan matriks default", "Masuk akal / terlalu merata / condong sipil"],
        ],
    )
    note(
        doc,
        "Yang tidak diubah live di peta.",
        "A1, A2, C2, E2 adalah keputusan arsitektur. Diskusi dan catat; jangan berjanji "
        "‘langsung kelihatan di peta hari ini’. Yang kelihatan live: slider jenjang, "
        "akreditasi, disiplin, pusat, kepakaran, spillover.",
    )

    heading(doc, "8. Kalibrasi peta (fasilitator)")
    body(doc, "Mode awal: Keselarasan, bahaya Gempabumi. Reset dulu jika slider sempat bergeser.")
    bullets(
        doc,
        [
            "Geser Teknik Sipil turun 0,20: apakah kuadran Senjang bergeser masuk akal?",
            "Naikkan S3, turunkan S1: apakah Jawa vs luar Jawa berubah wajar?",
            "Spillover ke 0: apakah Papua / NTT lebih ‘gelap’ pendidikan?",
            "Tanyakan: ‘Apakah pergeseran ini sesuai intuisi Bapak/Ibu sebagai ahli daerah?’",
            "Jangan reset di tengah kalimat; reset hanya jika peserta bingung.",
        ],
        numbered=True,
    )

    heading(doc, "9. Cadangan jika macet")
    table(
        doc,
        ["Masalah", "Tindakan"],
        [
            [
                "Form tidak kirim / error",
                "Lanjut diskusi peta. Catat A1–E2 di kertas atau Google Form darurat. Coba form lagi di jeda.",
            ],
            [
                "Wi-Fi lemah",
                "Hotspot. Atau satu laptop bersama untuk Kirim isian. Remote: minta isi nanti di hotel/wifi.",
            ],
            [
                "Peserta mencari tombol unduh",
                "Ulangi: ‘Kirim isian saja. JSON diunduh fasilitator.’",
            ],
            [
                "Pertanyaan inventaris PT",
                "Tunda. ‘Hari ini bobot, bukan kelengkapan daftar. Template data ada di halaman metodologi.’",
            ],
            [
                "Pertanyaan IRBI 2024 vs 2025",
                "Versi beku: IRBI 2025. Komposit resmi provinsi; per bahaya rata-rata kabupaten.",
            ],
        ],
    )

    heading(doc, "10. Setelah acara")
    bullets(
        doc,
        [
            "Unduh semua JSON dari /validasi/hasil (token 12345678).",
            "Simpan berkas di Drive terpisah (boleh ada nama); jangan commit data pribadi ke GitHub tanpa izin.",
            "Rekap: modus A1–E2; median sel matriks yang diisi; kutipan G1/G2 tanpa nama.",
            "Kirim ringkasan 1 halaman ke validator yang minta.",
            "Jika ada konsensus angka: ubah src/lib/weights.ts, metodologi, lalu deploy.",
        ],
        numbered=True,
    )

    heading(doc, "11. Logistik hibrid")
    bullets(
        doc,
        [
            "Onsite: QR di layar kedua atau cetak A5 di pintu.",
            "Remote: kirim tautan form di chat Zoom sekali di menit 10, sekali di menit 45.",
            "Jangan minta semua orang mengisi 64 sel matriks.",
            "Satu layar kebenaran: peta yang Anda geser, bukan 15 orang menggeser masing-masing.",
        ],
    )

    heading(doc, "Lampiran. Isi form (untuk diingat)")
    table(
        doc,
        ["Bagian", "Isi"],
        [
            ["1 Identitas", "Nama/inisial, institusi, keahlian, pernah lihat peta"],
            ["2 Angka risiko", "A1, alasan, A2"],
            ["3 Jenjang", "S1/D4, S2, S3 (kosong = 1,00), B1"],
            ["4 Akreditasi", "Internasional 1,00 … Baik 0,70; C1 IABEE; C2 tanpa peringkat"],
            ["5 Pusat & kepakaran", "D1 PkM, D2 provinsi tanpa pusat"],
            ["6 Spillover", "E1, E2"],
            ["7 Matriks", "F1 + sel yang diubah saja"],
            ["8 Catatan", "G1 ubah, G2 jangan ubah, G3 lain"],
            ["9 Persetujuan", "Bukan kesiapan; boleh setel bobot; nama institusi opsional"],
        ],
    )

    body(
        doc,
        "P2MI Multidisiplin FTSL ITB 2026 — Mainstreaming Disaster Resiliency in Infrastructure Systems.",
        italic=True,
        size=10,
    )

    OUT_PUBLIC.parent.mkdir(parents=True, exist_ok=True)
    OUT_DOCS.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PUBLIC)
    doc.save(OUT_DOCS)
    print(f"wrote {OUT_PUBLIC} ({OUT_PUBLIC.stat().st_size} bytes)")
    print(f"wrote {OUT_DOCS} ({OUT_DOCS.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
